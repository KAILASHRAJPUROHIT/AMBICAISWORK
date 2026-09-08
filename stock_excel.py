"""Read-only ingestion for dated stock workbooks.

This module is intentionally not wired into the live app yet.  It discovers
workbooks named DDMMYYYY.xls(x), validates the real stock-report schema, and
returns immutable records for later pipeline and dashboard integration.
Nothing here writes to a workbook or changes the Stock directory.
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Sequence


BASE = os.path.dirname(os.path.abspath(__file__))
STOCK_DIR = os.path.normpath(
    os.path.abspath(
        os.path.expandvars(
            os.path.expanduser(
                os.environ.get("AJ_STOCK_DIR", os.path.join(BASE, "Stock"))
            )
        )
    )
)

_DATED_WORKBOOK_RE = re.compile(
    r"^(?P<stamp>\d{8})\.(?P<extension>xls|xlsx)$",
    re.IGNORECASE,
)
_REPORT_DATE_RE = re.compile(
    r"\bAs\s+On\s+Date\s*:\s*(\d{1,2}/\d{1,2}/\d{4})\b",
    re.IGNORECASE,
)
_INVALID_FILENAME_CHARS_RE = re.compile(r'[\\/:*?"<>|]')

_HEADER_ALIASES = {
    "label_no": {"labelno", "labelnumber", "labeltagno", "tagno", "tagnumber"},
    "old_barcode_no": {"oldbarcodeno", "oldbarcodenumber", "oldbarcode"},
    "prefix": {"prefix"},
    "carat": {"carat", "karat"},
    "variety_name": {"varietyname", "variety"},
    "gross_weight": {"grosswt", "grossweight"},
    "net_weight": {"netwt", "netweight"},
    "pieces": {"pcs", "pieces", "piececount"},
    "huid": {"huid", "huids"},
}
_REQUIRED_HEADERS = tuple(_HEADER_ALIASES)
GENERAL_VARIETY = "GENERAL"


class StockWorkbookError(RuntimeError):
    """Base class for stock discovery, schema, and data errors."""


class StockSchemaError(StockWorkbookError):
    """The workbook does not contain the expected stock-report columns."""


class StockDataError(StockWorkbookError):
    """A stock row is incomplete, duplicated, or internally inconsistent."""


@dataclass(frozen=True, slots=True)
class StockRecord:
    label_no: str
    old_barcode_no: str
    prefix: str
    carat: str
    variety_name: str | None
    gross_weight: Decimal
    net_weight: Decimal
    pieces: int
    huids: tuple[str, ...]
    source_row: int

    @property
    def routing_variety(self) -> str:
        """Final output-folder key; blank source values use ``GENERAL``."""

        return self.variety_name or GENERAL_VARIETY

    def routing_key(self, ornament_type: str) -> tuple[str, str]:
        """Return ``(ornament_type, output variety)`` for pipeline routing."""

        normalised_ornament = ornament_type.strip()
        if not normalised_ornament:
            raise ValueError("ornament_type is required for stock routing")
        return normalised_ornament, self.routing_variety


@dataclass(frozen=True, slots=True)
class StockSnapshot:
    source_path: Path
    filename_date: date | None
    report_date: date | None
    sheet_name: str
    records: tuple[StockRecord, ...]
    source_total_gross_weight: Decimal | None
    source_total_net_weight: Decimal | None
    source_total_pieces: int | None

    @property
    def record_count(self) -> int:
        return len(self.records)

    @property
    def total_pieces(self) -> int:
        return sum(record.pieces for record in self.records)

    @property
    def varieties(self) -> tuple[str, ...]:
        return tuple(
            sorted(
                {
                    record.variety_name
                    for record in self.records
                    if record.variety_name is not None
                }
            )
        )

    @property
    def blank_variety_labels(self) -> tuple[str, ...]:
        return tuple(
            record.label_no
            for record in self.records
            if record.variety_name is None
        )

    def records_by_label(self) -> dict[str, StockRecord]:
        """Return a case-insensitive tag lookup keyed by ``label.casefold()``."""

        return {record.label_no.casefold(): record for record in self.records}

    def records_by_filename_label(self) -> dict[str, StockRecord]:
        """Return records keyed by the capture-safe filename stem.

        Capture replaces Windows-invalid characters such as ``/`` with ``_``.
        A collision is an ambiguous identity and must stop routing rather than
        silently choose one stock row.
        """

        indexed: dict[str, StockRecord] = {}
        for record in self.records:
            key = safe_filename_label(record.label_no).casefold()
            prior = indexed.get(key)
            if prior is not None and prior.label_no.casefold() != record.label_no.casefold():
                raise StockDataError(
                    f"Stock labels {prior.label_no!r} and {record.label_no!r} "
                    f"both map to capture filename {key!r}"
                )
            indexed[key] = record
        return indexed

    def lookup_label(self, label: str) -> StockRecord | None:
        """Look up either an original stock tag or a capture-safe filename."""

        key = label.strip().casefold()
        if not key:
            return None
        exact = self.records_by_label().get(key)
        if exact is not None:
            return exact
        return self.records_by_filename_label().get(key)


@dataclass(frozen=True, slots=True)
class StockLabelInventory:
    """Minimal current-stock identity list used for dashboard counts.

    Some daily exports contain only ``ItemName`` and ``Label No``.  They are
    authoritative for which tags are currently in stock, but must never be
    treated as routing records because they do not contain variety or weight
    data.
    """

    source_path: Path
    labels: tuple[str, ...]

    @property
    def record_count(self) -> int:
        return len(self.labels)


def safe_filename_label(label: str) -> str:
    """Mirror capture's authoritative Windows-safe tag filename transform."""

    return _INVALID_FILENAME_CHARS_RE.sub("_", label.strip())


def stock_date_from_filename(path: str | os.PathLike[str]) -> date | None:
    """Return the DDMMYYYY date encoded in a supported workbook filename."""

    match = _DATED_WORKBOOK_RE.fullmatch(Path(path).name)
    if match is None:
        return None
    try:
        return datetime.strptime(match.group("stamp"), "%d%m%Y").date()
    except ValueError:
        return None


def discover_stock_workbooks(
    stock_dir: str | os.PathLike[str] = STOCK_DIR,
) -> tuple[Path, ...]:
    """Return valid dated workbooks, oldest to newest."""

    root = Path(stock_dir)
    if not root.is_dir():
        raise FileNotFoundError(f"Stock directory does not exist: {root}")
    dated = [
        (stock_date_from_filename(path), path)
        for path in root.iterdir()
        if path.is_file() and stock_date_from_filename(path) is not None
    ]
    return tuple(path for _, path in sorted(dated, key=lambda item: (item[0], item[1].name)))


def latest_stock_workbook(
    stock_dir: str | os.PathLike[str] = STOCK_DIR,
    *,
    as_of: date | None = None,
) -> Path:
    """Select the newest dated workbook, failing on same-day ambiguity."""

    candidates = [
        path
        for path in discover_stock_workbooks(stock_dir)
        if as_of is None or stock_date_from_filename(path) <= as_of
    ]
    if not candidates:
        qualifier = f" on or before {as_of.isoformat()}" if as_of else ""
        raise FileNotFoundError(f"No dated stock workbook found{qualifier}")

    latest_date = stock_date_from_filename(candidates[-1])
    latest = [
        path for path in candidates if stock_date_from_filename(path) == latest_date
    ]
    if len(latest) != 1:
        names = ", ".join(path.name for path in latest)
        raise StockWorkbookError(
            f"Multiple stock workbooks exist for {latest_date.isoformat()}: {names}"
        )
    return latest[0]


def scheduled_scan_times(day: date) -> tuple[time, ...]:
    """Return the required local scan times for a calendar day."""

    scans = [time(hour=11)]
    if day.weekday() == 3:  # Thursday
        scans.append(time(hour=13))
    return tuple(scans)


def load_stock_workbook(path: str | os.PathLike[str]) -> StockSnapshot:
    """Read and validate one supported stock workbook without modifying it."""

    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(f"Stock workbook does not exist: {source}")

    extension = source.suffix.casefold()
    if extension == ".xls":
        sheets = _read_xls(source)
    elif extension == ".xlsx":
        sheets = _read_xlsx(source)
    else:
        raise StockWorkbookError(f"Unsupported stock workbook type: {source.suffix}")

    schema_errors: list[str] = []
    for sheet_name, rows in sheets:
        try:
            return _parse_sheet(source, sheet_name, rows)
        except StockSchemaError as exc:
            schema_errors.append(f"{sheet_name}: {exc}")

    details = "; ".join(schema_errors) if schema_errors else "workbook has no sheets"
    raise StockSchemaError(f"No stock-report sheet found in {source.name}: {details}")


def load_stock_label_inventory(
    path: str | os.PathLike[str],
) -> StockLabelInventory:
    """Read the unique labels from either a compact or rich stock export."""

    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(f"Stock workbook does not exist: {source}")
    extension = source.suffix.casefold()
    if extension == ".xls":
        sheets = _read_xls(source)
    elif extension == ".xlsx":
        sheets = _read_xlsx(source)
    else:
        raise StockWorkbookError(f"Unsupported stock workbook type: {source.suffix}")

    label_aliases = _HEADER_ALIASES["label_no"]
    for _sheet_name, rows in sheets:
        header_row = None
        label_column = None
        for row_index, row in enumerate(rows[:100]):
            for column_index, value in enumerate(row):
                if _normalise_header(value) in label_aliases:
                    header_row = row_index
                    label_column = column_index
                    break
            if header_row is not None:
                break
        if header_row is None or label_column is None:
            continue

        labels: list[str] = []
        seen: dict[str, int] = {}
        for zero_based_row, row in enumerate(
            rows[header_row + 1 :], start=header_row + 1
        ):
            label = _text(row[label_column] if label_column < len(row) else None)
            if (
                not label
                or label.casefold() in {"total", "grand total"}
                or _is_report_filter_label(label)
            ):
                continue
            key = label.casefold()
            if key in seen:
                raise StockDataError(
                    f"Duplicate Label No {label!r} at rows "
                    f"{seen[key]} and {zero_based_row + 1}"
                )
            seen[key] = zero_based_row + 1
            labels.append(label)
        if labels:
            return StockLabelInventory(source, tuple(labels))

    raise StockSchemaError(f"No Label No column found in {source.name}")


def load_stock_label_categories(
    path: str | os.PathLike[str],
) -> dict[str, str]:
    """Map current stock tag spellings and safe filename stems to ItemName.

    The newest daily stock workbook is authoritative. Both ``LR22/13`` and
    its capture filename form ``LR22_13`` resolve to the same category label.
    Ambiguous safe-filename collisions fail closed.
    """
    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(f"Stock workbook does not exist: {source}")
    extension = source.suffix.casefold()
    if extension == ".xls":
        sheets = _read_xls(source)
    elif extension == ".xlsx":
        sheets = _read_xlsx(source)
    else:
        raise StockWorkbookError(f"Unsupported stock workbook type: {source.suffix}")

    label_aliases = _HEADER_ALIASES["label_no"]
    item_aliases = {"itemname", "item", "ornamenttype", "category"}
    for _sheet_name, rows in sheets:
        header_row = label_column = item_column = None
        for row_index, row in enumerate(rows[:100]):
            normalised = [_normalise_header(value) for value in row]
            for column_index, value in enumerate(normalised):
                if value in label_aliases:
                    label_column = column_index
                if value in item_aliases:
                    item_column = column_index
            if label_column is not None and item_column is not None:
                header_row = row_index
                break
        if header_row is None:
            continue

        result: dict[str, str] = {}
        owners: dict[str, str] = {}
        for zero_based_row, row in enumerate(rows[header_row + 1 :], start=header_row + 1):
            label = _text(row[label_column] if label_column < len(row) else None)
            item_name = _text(row[item_column] if item_column < len(row) else None)
            if not label or label.casefold() in {"total", "grand total"}:
                continue
            # Ornate occasionally appends report predicates, e.g.
            # ``[Carat] <> 'S925'``, below the totals row. They are not stock
            # tags and have no category. Ignore only this explicit metadata;
            # all actual Label No + ItemName pairs remain authoritative.
            if _is_report_filter_label(label) and not item_name:
                continue
            if not item_name:
                raise StockDataError(
                    f"{source.name} row {zero_based_row + 1} has Label No {label!r} but no ItemName"
                )
            for key in {label.casefold(), safe_filename_label(label).casefold()}:
                prior_label = owners.get(key)
                prior_item = result.get(key)
                if prior_label is not None and (prior_label.casefold() != label.casefold() or prior_item != item_name):
                    raise StockDataError(
                        f"Stock labels {prior_label!r} and {label!r} collide at filename key {key!r}"
                    )
                owners[key] = label
                result[key] = item_name
        if result:
            return result

    raise StockSchemaError(f"No ItemName + Label No columns found in {source.name}")


def _is_report_filter_label(value: str) -> bool:
    """Ornate appends bracketed report predicates below stock rows."""
    return value.startswith("[") and "]" in value


def _read_xls(path: Path) -> list[tuple[str, list[list[Any]]]]:
    try:
        import xlrd
    except ImportError as exc:
        raise StockWorkbookError(
            "Reading .xls stock files requires xlrd>=2.0.1"
        ) from exc

    try:
        workbook = xlrd.open_workbook(str(path), on_demand=True)
    except Exception as exc:
        raise StockWorkbookError(f"Could not read {path.name}: {exc}") from exc

    sheets: list[tuple[str, list[list[Any]]]] = []
    try:
        for sheet_name in workbook.sheet_names():
            sheet = workbook.sheet_by_name(sheet_name)
            rows = [
                [sheet.cell_value(row, column) for column in range(sheet.ncols)]
                for row in range(sheet.nrows)
            ]
            sheets.append((sheet_name, rows))
    finally:
        workbook.release_resources()
    return sheets


def _read_xlsx(path: Path) -> list[tuple[str, list[list[Any]]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise StockWorkbookError(
            "Reading .xlsx stock files requires openpyxl>=3.1"
        ) from exc

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise StockWorkbookError(f"Could not read {path.name}: {exc}") from exc

    try:
        return [
            (sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)])
            for sheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def _normalise_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", _text(value).casefold())


def _find_header(rows: Sequence[Sequence[Any]]) -> tuple[int, dict[str, int]]:
    aliases_to_field = {
        alias: field
        for field, aliases in _HEADER_ALIASES.items()
        for alias in aliases
    }
    for row_index, row in enumerate(rows[:100]):
        columns: dict[str, int] = {}
        for column_index, value in enumerate(row):
            field = aliases_to_field.get(_normalise_header(value))
            if field is not None:
                columns[field] = column_index
        if all(field in columns for field in _REQUIRED_HEADERS):
            return row_index, columns
    raise StockSchemaError(
        "required columns not found: " + ", ".join(_REQUIRED_HEADERS)
    )


def _parse_sheet(
    source: Path,
    sheet_name: str,
    rows: Sequence[Sequence[Any]],
) -> StockSnapshot:
    header_row, columns = _find_header(rows)
    report_date = _find_report_date(rows[:header_row])
    records: list[StockRecord] = []
    seen_labels: dict[str, int] = {}
    source_totals: tuple[Decimal, Decimal, int] | None = None

    for zero_based_row, row in enumerate(rows[header_row + 1 :], start=header_row + 1):
        values = {
            field: row[column] if column < len(row) else None
            for field, column in columns.items()
        }
        if not any(_text(value) for value in values.values()):
            continue

        label_no = _text(values["label_no"])
        if not label_no:
            total = _summary_totals(values, zero_based_row + 1)
            if total is None:
                raise StockDataError(
                    f"{sheet_name} row {zero_based_row + 1} has data but no Label No"
                )
            source_totals = total
            continue
        if _normalise_header(label_no) in _HEADER_ALIASES["label_no"]:
            continue
        if label_no.casefold() in {"total", "grand total"}:
            source_totals = _required_totals(values, zero_based_row + 1)
            continue

        record = StockRecord(
            label_no=label_no,
            old_barcode_no=_required_text(
                values["old_barcode_no"], "Old BarcodeNo", zero_based_row + 1
            ),
            prefix=_required_text(values["prefix"], "Prefix", zero_based_row + 1),
            carat=_required_text(values["carat"], "Carat", zero_based_row + 1),
            variety_name=_text(values["variety_name"]) or None,
            gross_weight=_required_decimal(
                values["gross_weight"], "Gross Wt", zero_based_row + 1
            ),
            net_weight=_required_decimal(
                values["net_weight"], "Net Wt", zero_based_row + 1
            ),
            pieces=_required_integer(values["pieces"], "Pcs", zero_based_row + 1),
            huids=tuple(
                part.strip()
                for part in _text(values["huid"]).split(",")
                if part.strip()
            ),
            source_row=zero_based_row + 1,
        )
        lookup_key = record.label_no.casefold()
        if lookup_key in seen_labels:
            raise StockDataError(
                f"Duplicate Label No {record.label_no!r} at rows "
                f"{seen_labels[lookup_key]} and {record.source_row}"
            )
        seen_labels[lookup_key] = record.source_row
        records.append(record)

    if not records:
        raise StockDataError(f"{sheet_name} contains no stock records")

    snapshot = StockSnapshot(
        source_path=source,
        filename_date=stock_date_from_filename(source),
        report_date=report_date,
        sheet_name=sheet_name,
        records=tuple(records),
        source_total_gross_weight=source_totals[0] if source_totals else None,
        source_total_net_weight=source_totals[1] if source_totals else None,
        source_total_pieces=source_totals[2] if source_totals else None,
    )
    _validate_source_totals(snapshot)
    return snapshot


def _find_report_date(rows: Iterable[Sequence[Any]]) -> date | None:
    for row in rows:
        for value in row:
            match = _REPORT_DATE_RE.search(_text(value))
            if match:
                try:
                    return datetime.strptime(match.group(1), "%d/%m/%Y").date()
                except ValueError as exc:
                    raise StockDataError(
                        f"Invalid report date {match.group(1)!r}"
                    ) from exc
    return None


def _summary_totals(
    values: dict[str, Any],
    row_number: int,
) -> tuple[Decimal, Decimal, int] | None:
    if not any(
        _text(values[field]) for field in ("gross_weight", "net_weight", "pieces")
    ):
        return None
    return _required_totals(values, row_number)


def _required_totals(
    values: dict[str, Any],
    row_number: int,
) -> tuple[Decimal, Decimal, int]:
    return (
        _required_decimal(values["gross_weight"], "total Gross Wt", row_number),
        _required_decimal(values["net_weight"], "total Net Wt", row_number),
        _required_integer(values["pieces"], "total Pcs", row_number),
    )


def _validate_source_totals(snapshot: StockSnapshot) -> None:
    if snapshot.source_total_pieces is None:
        return
    computed_gross = sum(
        (record.gross_weight for record in snapshot.records),
        start=Decimal("0"),
    )
    computed_net = sum(
        (record.net_weight for record in snapshot.records),
        start=Decimal("0"),
    )
    if (
        computed_gross != snapshot.source_total_gross_weight
        or computed_net != snapshot.source_total_net_weight
        or snapshot.total_pieces != snapshot.source_total_pieces
    ):
        raise StockDataError(
            "Workbook totals do not match its stock rows: "
            f"gross {computed_gross}/{snapshot.source_total_gross_weight}, "
            f"net {computed_net}/{snapshot.source_total_net_weight}, "
            f"pieces {snapshot.total_pieces}/{snapshot.source_total_pieces}"
        )


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _required_text(value: Any, field: str, row_number: int) -> str:
    text_value = _text(value)
    if not text_value:
        raise StockDataError(f"Row {row_number} has blank {field}")
    return text_value


def _required_decimal(value: Any, field: str, row_number: int) -> Decimal:
    text_value = _text(value)
    if not text_value:
        raise StockDataError(f"Row {row_number} has blank {field}")
    try:
        return Decimal(text_value)
    except InvalidOperation as exc:
        raise StockDataError(
            f"Row {row_number} has invalid {field}: {text_value!r}"
        ) from exc


def _required_integer(value: Any, field: str, row_number: int) -> int:
    number = _required_decimal(value, field, row_number)
    if number != number.to_integral_value():
        raise StockDataError(
            f"Row {row_number} has non-integer {field}: {number}"
        )
    return int(number)
